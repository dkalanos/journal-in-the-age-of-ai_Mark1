"""Build system_diagrams v6dk.pptx from a fresh copy of v5dk.

Run:  python figures_v6.py
Then: powershell -NoProfile -File export_figures.ps1
      (PowerPoint; writes PNG and PDF per slide into this folder)

What changes from v5dk (Kalanoski, 2026-10-29), both on Fig. 3 (slide 4):
- The "New" box names the four kinds of newness decided on 2026-10-28 (answer, question, reading, evidence) in
  place of the October 17 kinds, and says they are judged on the question the paper states (protocol formal-0.4,
  Mark 1 CHANGES item 41).
- The Level 2 outcome box follows mapping 0.3 (CHANGES item 43): return for corrections only where a pointer does
  not resolve or the evidence says nothing; evidence that carries part of a claim travels to review with the
  finding. The v5 wording returned every Level 2 finding.
Everything else, the record/package wording included, is as in v5dk. Shape ids are v5dk's.
"""
import shutil
from pathlib import Path

from pptx import Presentation

from figures_v3 import set_text

HERE = Path(__file__).parent
SRC = HERE / "system_diagrams v5dk.pptx"
DECK = HERE / "system_diagrams v6dk.pptx"


def fig3(prs):
    s = prs.slides[3]
    set_text(s, 12, [[("New: ", True, False),
                      ("does the paper add to what is known: an answer to a question already posed, a question "
                       "not posed before, a reading of results in hand, or new evidence on a claim already "
                       "stated? Judged on the question the paper states.", False, False)],
                     [("Interesting: ", True, False),
                      ("should this community attend to it now (Davis 1971)?", False, False)]], 8.5)
    set_text(s, 27, ["Unresolved pointer or empty evidence: return for corrections, no new work (match paper, "
                     "appendix, record). Partial evidence goes to review with the finding."], 8)


def main():
    shutil.copyfile(SRC, DECK)  # idempotent rebuild from v5dk
    prs = Presentation(DECK)
    fig3(prs)
    prs.save(DECK)
    print("saved", DECK, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
