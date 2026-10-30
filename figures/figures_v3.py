"""Build the v3 figures from notebook/vision/system-design.md into system_diagrams v3mdr.pptx.

Run with the project venv:  ~/.venvs/drt/bin/python figures_v3.py
Rebuilds the deck from a fresh copy of v2claude; the v1/v2claude decks are not touched.
Slide order after the build: legend, Fig. 1 (before/after), Fig. 2 (record), Fig. 3 (levels),
Fig. 4 (one paper's path), Fig. 4a, Fig. 4b (algorithmic maps).
Styling follows Dimitrija's legend: role colour = theme accent (1 journal, 2 author+AI,
3 machine, 4 named people), tinted fill + solid 1pt border = built/specified, Times New Roman.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
DECK = HERE / "system_diagrams v3mdr.pptx"
FONT = "Times New Roman"

ROLE = {
    "journal": MSO_THEME_COLOR.ACCENT_1,
    "author": MSO_THEME_COLOR.ACCENT_2,
    "machine": MSO_THEME_COLOR.ACCENT_3,
    "people": MSO_THEME_COLOR.ACCENT_4,
    "tag": MSO_THEME_COLOR.ACCENT_5,
}
GREY_FILL = RGBColor(0xEE, 0xEF, 0xF1)
GREY_LINE = RGBColor(0x9A, 0x9E, 0xA6)
TEXT = RGBColor(0x1E, 0x1E, 0x24)
MUTED = RGBColor(0x5C, 0x60, 0x68)


def _runs(tf, parts, size, align=PP_ALIGN.LEFT, color=TEXT):
    """parts: list of paragraphs; each paragraph a list of (text, bold, italic) or a str."""
    tf.word_wrap = True
    first = True
    for para in parts:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if isinstance(para, str):
            para = [(para, False, False)]
        for text, bold, italic in para:
            r = p.add_run()
            r.text = text
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color


def box(slide, x, y, w, h, parts, *, role=None, fill="tint", size=9, shape=MSO_SHAPE.RECTANGLE,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, dash=None, color=TEXT):
    """A styled box. fill: 'tint' (role colour at 78% lighter), 'solid' (role colour, white
    text), 'white', 'grey', or None (no fill)."""
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    elif fill == "grey":
        sh.fill.solid(); sh.fill.fore_color.rgb = GREY_FILL
    elif fill == "white":
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    elif fill == "solid":
        sh.fill.solid(); sh.fill.fore_color.theme_color = ROLE[role]
        color = RGBColor(0xFF, 0xFF, 0xFF)
    else:
        sh.fill.solid(); sh.fill.fore_color.theme_color = ROLE[role]
        sh.fill.fore_color.brightness = 0.78
    if role is None and fill in ("grey", "white", None):
        sh.line.color.rgb = GREY_LINE
    else:
        sh.line.color.theme_color = ROLE[role]
    sh.line.width = Pt(1)
    if dash:
        sh.line.dash_style = dash
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.07)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    _runs(tf, parts, size, align=align, color=color)
    return sh


def text(slide, x, y, w, h, parts, *, size=9, align=PP_ALIGN.LEFT, color=TEXT,
         anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    _runs(tf, parts, size, align=align, color=color)
    return tb


def title(slide, fig, heading, subtitle):
    text(slide, 0.4, 0.25, 12.5, 0.5, [[(f"{fig} - {heading}", True, False)]], size=24)
    text(slide, 0.4, 0.78, 12.5, 0.4, [[(subtitle, False, True)]], size=11, color=MUTED)


def legend(slide, y=6.95):
    items = [("journal", "Journal"), ("author", "Author + AI"), ("machine", "Machine"),
             ("people", "Named people")]
    x = 0.4
    for role, label in items:
        box(slide, x, y, 0.22, 0.18, [""], role=role, fill="tint")
        text(slide, x + 0.26, y - 0.03, 1.2, 0.25, [label], size=8.5)
        x += 1.45
    box(slide, x, y, 0.22, 0.18, [""], fill="grey")
    text(slide, x + 0.26, y - 0.03, 1.6, 0.25, ["today's arrangement"], size=8.5)


# ---------------------------------------------------------------------------------------
# Fig. 1 — The journal, before and after
# ---------------------------------------------------------------------------------------
ROWS = [
    # (row label, today, the cost that fell, after, role of the after cell)
    ("The journal publishes",
     "Formatting instructions and a statement of scope.",
     "",
     [[("Five objects, versioned: ", True, False),
       ("the record protocol (shared across journals), the checks, the standard they "
        "apply, the mapping from findings to actions, and this community's interest "
        "criteria.", False, False)]],
     "journal"),
    ("The author submits",
     "A manuscript. The evidence stays with the author and is taken on trust; a replication "
     "package follows acceptance, if ever.",
     "keeping a complete record",
     [[("The research record: ", True, False),
       ("narrative, AI appendix, materials, provenance, declarations, kept by AI as the work "
        "happens. Working history at the author's choice.", False, False)]],
     "author"),
    ("The paper is",
     "One linear narrative, fixed length, compressed methods, a reference list standing in "
     "for the sources.",
     "pages (the internet)",
     [[("Two layers: ", True, False),
       ("a narrative of any length, for people; an AI appendix, for the reader that evaluates. "
        "Their correspondence is checkable.", False, False)]],
     "author"),
    ("Who verifies",
     "Reviewers, unpaid, by reading. Most of what a paper rests on is never checked.",
     "inspecting and re-executing a record",
     [[("The machine, at three levels: ", True, False),
       ("form and correspondence (integrity); warrant against the standard, the literature, "
        "and re-execution. Findings with pointers, never a verdict. Authors run the same checks "
        "first.", False, False)]],
     "machine"),
    ("Who decides, and how",
     "A desk, in private: a verdict with no stated rule and no name attached.",
     "reading at scale",
     [[("Four signed decisions ", True, False),
       ("by named people: intake by a published mapping; new, on the evidence map; interest, "
        "under published criteria; certify. Departures recorded; their rate published.",
        False, False)]],
     "people"),
    ("Reviewers",
     "Verify and judge at once, re-reading the whole paper every round.",
     "",
     [[("Judge only. ", True, False),
       ("They receive the narrative, the findings report, and the evidence map, and are not "
        "asked to verify.", False, False)]],
     "people"),
    ("The journal certifies",
     "“Published in Journal X”: one undifferentiated signal of quality and importance.",
     "",
     [[("One certificate, or nothing: ", True, False),
       ("meets standard vX; new, in this kind; of interest to this community; signed by these "
        "three named people.", False, False)]],
     "journal"),
    ("Between journals",
     "Nothing travels, and every rejection is a private anchor on the next reading.",
     "",
     [[("The record travels, always. ", True, False),
       ("Everything else only at the author's choice.", False, False)]],
     "journal"),
]


def fig1_before_after(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    title(s, "Fig. 1", "The journal, before and after",
          "The same institution, with the print-era arrangements replaced by arrangements "
          "for a world in which keeping, inspecting, and reading records at scale are cheap.")
    # column geometry (inches)
    x_label, w_label = 0.4, 1.55
    x_today, w_today = 2.0, 3.9
    x_cost, w_cost = 6.0, 1.55
    x_after, w_after = 7.65, 5.3
    y0, row_h, gap = 1.55, 0.6, 0.06
    # column headers
    text(s, x_today, 1.2, w_today, 0.3, [[("Today — arrangements for print-era costs", True, False)]],
         size=10, color=MUTED)
    text(s, x_cost, 1.2, w_cost, 0.3, [[("The cost that fell", True, False)]], size=10,
         color=MUTED, align=PP_ALIGN.CENTER)
    text(s, x_after, 1.2, w_after, 0.3, [[("After — the same journal, redesigned", True, False)]],
         size=10, color=MUTED)
    y = y0
    for label, today, cost, after, role in ROWS:
        text(s, x_label, y, w_label, row_h, [[(label, True, False)]], size=9.5,
             anchor=MSO_ANCHOR.MIDDLE)
        box(s, x_today, y, w_today, row_h, [today], fill="grey", size=8.5, anchor=MSO_ANCHOR.MIDDLE)
        if cost:
            text(s, x_cost, y, w_cost, row_h, [[(cost, False, True)]], size=8.5, color=MUTED,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x_cost + 0.45),
                                       Inches(y + row_h - 0.16), Inches(0.65), Inches(0.1))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = GREY_LINE
            arrow.line.fill.background()
        else:
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x_cost + 0.45),
                                       Inches(y + row_h / 2 - 0.05), Inches(0.65), Inches(0.1))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = GREY_LINE
            arrow.line.fill.background()
        box(s, x_after, y, w_after, row_h, after, role=role, fill="tint", size=8.5,
            anchor=MSO_ANCHOR.MIDDLE)
        y += row_h + gap
    # what stays the same
    text(s, 0.4, y + 0.02, 12.5, 0.4,
         [[("What does not change: ", True, False),
           ("the author works however they like; what a community finds interesting is decided "
            "by its members; people read the paper.", False, False)]], size=9)
    legend(s, y=7.05)
    return s


# ---------------------------------------------------------------------------------------
# In-place edits to Dimitrija's slides
# ---------------------------------------------------------------------------------------
def shape_by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_text(slide, sid, parts, size, *, align=None, color=TEXT):
    """Replace a shape's text, keeping the shape. parts as in _runs."""
    sh = shape_by_id(slide, sid)
    tf = sh.text_frame
    old_align = tf.paragraphs[0].alignment if align is None else align
    # drop every paragraph but the first, then clear the first
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    _runs(tf, parts, size, align=old_align or PP_ALIGN.LEFT, color=color)


def delete(slide, *sids):
    for sid in sids:
        el = shape_by_id(slide, sid)._element
        el.getparent().remove(el)


def set_fill(slide, sid, kind, role=None):
    """Build-state fill: 'tint' = built for Nov 1 (role colour at 78% lighter);
    'white' = specified, not built. Border untouched."""
    sh = shape_by_id(slide, sid)
    if kind == "white":
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    else:
        sh.fill.solid(); sh.fill.fore_color.theme_color = ROLE[role]
        sh.fill.fore_color.brightness = 0.78


def lead(head, body):
    """One paragraph: bold lead-in, then body."""
    return [[(head, True, False), (body, False, False)]]


def fig2_path(prs):
    """Dimitrija's Fig. 1 (slide 2) → Fig. 4, one paper's path, per system-design.md."""
    s = prs.slides[1]
    set_text(s, 2, [[("Fig. 4 - The system: one paper's path", True, False)]], 24)
    set_text(s, 3, [[("A journal publishes its standard; AI keeps the record and checks it; "
                      "named people decide, and sign.", False, True)]], 11)
    # top row: before and during the work
    set_text(s, 4, lead("The journal publishes, versioned: ",
                        "the record protocol (a field spec shared across journals), the checks, "
                        "the standard the checks apply, the mapping from findings to actions, and "
                        "its community's interest criteria."), 10)
    set_text(s, 5, lead("The author works however they like. ",
                        "AI keeps the research record as the work proceeds, and the working "
                        "history accrues from the start. How the author meets the standard is the "
                        "author's business."), 10)
    set_text(s, 6, lead("The research record: ",
                        "the narrative and the AI appendix; data, code, materials; provenance "
                        "linking each result to what produced it; declarations of type, standard "
                        "version, and AI use."), 10)
    set_text(s, 7, lead("Submission: ",
                        "the record, with the author's own findings report from the published "
                        "checks. Working history at the author's choice. The journal receives "
                        "records the author has already checked, and compares its run to theirs."), 10)
    # bottom row: intake to certificate
    set_text(s, 9, lead("Intake checks, re-run by the journal: ",
                        "form and correspondence, internal to the record (integrity); warrant "
                        "against the standard, the literature, and re-execution (external). Output: "
                        "findings with pointers, and an evidence map. Never a verdict."), 10)
    set_text(s, 10, lead("Intake decision: ",
                         "a named editor applies the published mapping — proceed to review, "
                         "return for corrections, return without review, refer. Departures are "
                         "recorded with their reason. The editor then signs whether the "
                         "contribution is new, on the evidence map, after the author has answered "
                         "it once."), 10)
    set_text(s, 11, lead("Review and interest: ",
                         "reviewers receive the narrative, the findings report, and the map; they "
                         "judge, and do not verify. The editor and two named reviewers judge "
                         "whether what is new is of interest to this community, under the "
                         "published criteria."), 10)
    set_text(s, 13, lead("Certificate, or nothing: ",
                         "meets standard vX; new, in this kind; of interest to this community; "
                         "signed by these three named people. No partial certificate exists, and "
                         "nothing travels to another journal unless the author sends it."), 10)
    # edge labels
    set_text(s, 22, ["findings + evidence map"], 8.5)
    set_text(s, 26, ["certify"], 8.5)
    set_text(s, 28, ["return for corrections (record fixed, no new work), or without review "
                     "(strengthen the evidence or weaken the claims)"], 8.5)
    shape_by_id(s, 28).width = Inches(1.75)
    set_text(s, 30, ["revise and resubmit: new work; each round a new submission against the "
                     "same review history"], 8.5)
    # participation box
    set_text(s, 31, lead("Clean records travel by choice: ",
                         "a submission arrives with the author's own findings report or it does "
                         "not, and withholding it is read (Grossman 1981; Milgrom 1981). The "
                         "working history, submitted by choice, is evidence of human involvement "
                         "in the sense an authorship statement is."), 8)
    # limits row
    set_text(s, 32, lead("Limits, and what the design does about each.", ""), 10)
    set_text(s, 33, lead("Integrity, warrant, truth: ",
                         "an AI-written paper passes form and correspondence by construction; "
                         "warrant reaches outside the record; data that never existed is beyond "
                         "every level unless its source is public."), 9)
    set_text(s, 34, lead("Interest is human by authority, not capability: ",
                         "what a community finds interesting is a fact about the community "
                         "(Davis 1971). No advance in AI moves it."), 9.5)
    set_text(s, 35, lead("No pointer, no flag: ",
                         "every finding links to the record and to the standard applied, "
                         "versioned with the model that produced it, so it can be repeated and "
                         "challenged."), 9)
    set_text(s, 36, lead("The record cannot prove human direction: ",
                         "anything checkable could be generated. The AI-use declaration and the "
                         "working history rest on good faith and reputation, as authorship "
                         "statements do."), 9.5)
    set_text(s, 38, lead("Nothing travels without the author: ",
                         "each journal re-runs the checks and judges interest for its own "
                         "community; review history goes to another journal only at the author's "
                         "request."), 9.5)
    set_text(s, 40, lead("Deference: ",
                         "if editors never depart from the mapping, the mapping decides. The "
                         "published departure rate measures the editors; the guardrail cases "
                         "measure the mapping."), 9.5)
    # P tags and legend entry: the proposals are decided (system-design.md); dotted "proposed"
    # borders become solid "specified"
    delete(s, 8, 12, 37, 39, 55, 56, 57)
    for sid in (7, 11, 36, 38):
        shape_by_id(s, sid).line.dash_style = MSO_LINE.SOLID
    set_fill(s, 11, "white")  # review and interest: specified, not built
    # fit: the densest boxes
    for sid, size in ((9, 9.5), (10, 9), (11, 9.5), (36, 9)):
        sh = shape_by_id(s, sid)
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)


def fig3_levels(prs):
    """Dimitrija's Fig. 3 (slide 5): the levels table, per system-design.md Q1, Q5, Q6."""
    s = prs.slides[4]
    set_text(s, 2, [[("Fig. 3 - Levels of assessment: what is checked, against what, and who "
                      "decides", True, False)]], 24)
    set_text(s, 3, [[("Checks internal to the record verify integrity; checks that reach "
                      "outside it verify warrant. Findings at Levels 1–3; signed decisions "
                      "above the line.", False, True)]], 11)
    # Level 4
    set_text(s, 10, [[("4 · New, then interesting", True, False)]], 9.5)
    set_text(s, 12, [[("New: ", True, False),
                      ("does the paper add to what is known — a gap filled, a new insight, a "
                       "new stream, new evidence on an existing claim? ", False, False),
                      ("Interesting: ", True, False),
                      ("should this community attend to it now (Davis 1971)?", False, False)]], 9)
    set_text(s, 13, ["New: one editor, on the map. Interesting: the editor and two named "
                     "reviewers."], 9)
    set_text(s, 14, ["Two signed judgments under published criteria; then one certificate, "
                     "or nothing."], 9)
    set_text(s, 15, ["Certified, or returned with nothing travelling. No partial certificate."], 9)
    set_text(s, 37, ["Above the line, named people decide and sign. Below it, the machine emits "
                     "findings with pointers, never a verdict."], 9)
    # Level 3
    set_text(s, 18, ["Do the claims match the evidence offered, judged against what the record "
                     "does not supply: the published standard for the modality, the literature "
                     "(the evidence map), re-execution under perturbation? External: verifies "
                     "warrant."], 9)
    set_text(s, 19, ["Machine, reaching outside the record (SIEVE modules)"], 9)
    set_text(s, 20, ["Findings with a pointer to the record and to the standard applied; the "
                     "evidence map."], 9)
    set_text(s, 21, ["Strengthen the evidence or weaken the claims. Above the published "
                     "threshold: return without review."], 9)
    # Level 2
    set_text(s, 24, ["Does the narrative match the AI appendix, and the appendix the record? "
                     "Completeness, resolution, anchors. Internal: verifies integrity."], 9)
    set_text(s, 26, ["Findings checkable from the record alone."], 9)
    set_text(s, 27, ["Correct the narrative to match the record (or the record, then re-run). "
                     "Return for corrections; no new work."], 9)
    # Level 1
    set_text(s, 30, ["Does the record conform to the published protocol for its declared type: "
                     "parts present, declarations made, appendix well-formed? Internal."], 9)
    set_text(s, 32, ["Findings; nothing is rejected on form."], 9)
    set_text(s, 33, ["The author's AI brings the record to the protocol. Return for "
                     "corrections."], 9)
    # notes
    set_text(s, 38, ["Findings describe the paper; actions are what the journal does next. The "
                     "published mapping links them, a named editor applies it, and departures "
                     "are recorded with their reason. An AI-written paper passes Levels 1–2 "
                     "by construction; Level 3 is where the AI-age work is done."], 9.5)
    set_text(s, 40, lead("Verification — ",
                         "the machine checks the record, internally for integrity and externally "
                         "for warrant, links each finding to evidence and to the standard applied, "
                         "and versions it with the model that produced it."), 9.5)
    set_text(s, 41, lead("Judgment — ",
                         "named people decide whether the contribution is new and whether it is "
                         "of interest to this community, under published criteria, and sign."), 9.5)
    set_text(s, 42, lead("Accountability — ",
                         "people answer for decisions, not for findings. The mapping is published, "
                         "departures are recorded, and the departure rate is public, so a journal "
                         "that automates is visible."), 9.5)
    set_text(s, 47, [[("A desk rejection today: ", True, False),
                      ("private, unexplained, unattributed, uncontestable.", False, False)],
                     [("A return here: ", True, False),
                      ("a finding with a pointer, a published rule, and a name; the author holds "
                       "their own report from the same checks, so a return rests on a finding "
                       "they saw or on one marked as absent from their run.",
                       False, False)]], 9.5)
    delete(s, 11, 60, 61, 62)
    shape_by_id(s, 10).line.dash_style = MSO_LINE.SOLID
    set_fill(s, 10, "white")  # Level 4: specified, not built


def fig4_record(prs):
    """Dimitrija's Fig. 4 (slide 6) → Fig. 2: the record and the AI appendix, per Q3, Q4, Q5."""
    s = prs.slides[5]
    set_text(s, 2, [[("Fig. 2 - The research record and the AI appendix", True, False)]], 24)
    set_text(s, 3, [[("A machine can verify but reads narrative inefficiently; a person needs the "
                      "argument but cannot check the algebra. The paper has two layers, and their "
                      "correspondence is a checkable object.", False, True)]], 11)
    set_text(s, 5, [[("The research record: what the author submits", True, False)]], 11)
    # record parts, re-spaced to fit the seven-part appendix
    geom = {6: (1.59, 0.42), 7: (2.07, 1.05), 8: (3.18, 0.38), 9: (3.62, 0.72), 10: (4.40, 0.38)}
    for sid, (top, h) in geom.items():
        sh = shape_by_id(s, sid); sh.top = Inches(top); sh.height = Inches(h)
    set_text(s, 6, lead("Narrative: ", "for people. The question, what was done, what it means; "
                        "any length the argument needs."), 9)
    set_text(s, 7, lead("AI appendix: ", "for the reader that evaluates. Seven parts: definitions "
                        "of record; the procedure, re-executable; the claims list, each claim with "
                        "a pointer to its evidence; the authors' choices and why; reproduction "
                        "anchors with known values; scope and do-nots; declarations."), 9)
    set_text(s, 8, lead("Materials: ", "data, code, outputs, or field materials, each in a "
                        "declared access tier (public, restricted, described-only)."), 9)
    set_text(s, 9, lead("Provenance: ", "what produced each result — code and data (empirical); "
                        "the derivation (formal); seeds and runs (simulation); the audit trail from "
                        "observation to interpretation (qualitative); the argument structure "
                        "(verbal theory); corpus and inclusion criteria (review)."), 9)
    set_text(s, 10, lead("Declarations: ", "modality · standard version · access tier · how "
                         "AI was used. Working history at the author's choice."), 9)
    # checks
    set_text(s, 11, [[("Level 2 · Correspondence (internal)", True, False)],
                     [("Does the narrative match the appendix, and the appendix the record?",
                       False, False)],
                     [("(a) Completeness: ", True, False),
                      ("every claim the contribution rests on is in the claims list.", False, False)],
                     [("(b) Resolution: ", True, False),
                      ("each pointer resolves to evidence that says what the claim says.", False, False)],
                     [("(c) Anchors: ", True, False),
                      ("the reproduction anchors reproduce (deterministic).", False, False)],
                     [("Checkable from the record alone. Fix: correct the narrative to match the "
                       "record, or the record, then re-run.", False, True)]], 9.5)
    set_text(s, 12, lead("Diff between rounds (stretch): ",
                         "what changed in the record between review rounds, and whether the "
                         "evidence still supports the revised claims. What reviewers read from "
                         "the second round."), 9.5)
    set_text(s, 13, [[("Level 3 · Warrant (external)", True, False)],
                     [("Do the claims match the evidence offered, judged against what the record "
                       "does not supply?", False, False)],
                     [("The standard for the modality: ", True, False),
                      ("what a causal claim, a proposition, a qualitative interpretation needs "
                       "(SIEVE modules: statistical, internal, construct, external validity; "
                       "contribution; modality soundness).", False, False)],
                     [("The literature: ", True, False),
                      ("does the contribution claim survive the nearest prior work? Output: the "
                       "evidence map.", False, False)],
                     [("Re-execution under perturbation: ", True, False),
                      ("do the results survive the robustness checks the standard names?",
                       False, False)],
                     [("Fix: strengthen the evidence or weaken the claims.", False, True)]], 9.5)
    set_text(s, 14, [[("Slop is a claim that goes beyond its evidence. ", True, False),
                      ("It is not: (a) unconventional work — strong evidence, unfamiliar "
                       "approach; (b) wrong work — reasonable evidence, later overturned; "
                       "(c) unimportant work — well supported, little valued.", False, False)],
                     [("Internal checks cannot tell a Potemkin village from a town; Level 3 is "
                       "built to, and names what it cannot reach: data that never existed.",
                       False, False)]], 9.5)
    set_text(s, 19, lead("Modalities: ", "formal built for Nov 1; empirical in Stage 2; the "
                         "rest specified by protocol."), 10)
    set_text(s, 20, [[("Formal theory  built for Nov 1", True, False)]], 10, align=PP_ALIGN.CENTER)
    set_text(s, 21, [[("Empirical  Stage 2", True, False)]], 10, align=PP_ALIGN.CENTER)
    set_fill(s, 21, "white")
    set_text(s, 26, lead("Limit: the record cannot prove human direction. ",
                         "Anything checkable could be generated; the AI-use declaration and the "
                         "working history rest on good faith and reputation, as authorship "
                         "statements do."), 9.5)
    set_text(s, 28, lead("What is new, taken together: ",
                         "a claims list as a schema; narrative-to-record correspondence checked by "
                         "published code under a published standard; a published mapping from "
                         "findings to actions; one indivisible certificate signed by named people. "
                         "Registered Reports, Code Ocean, Statcheck, and eLife each did one of "
                         "these."), 9)
    set_text(s, 24, [[("Verbal theory", True, False)]], 10, align=PP_ALIGN.CENTER)
    delete(s, 27, 43, 44, 45)
    shape_by_id(s, 26).line.dash_style = MSO_LINE.SOLID


def fig4a_map(prs):
    """Dimitrija's Fig. 2a (slide 3) → Fig. 4a: published standard to findings report, per
    appendix C.3 / D. Boxes keep their geometry; text and borders change."""
    s = prs.slides[2]
    set_text(s, 2, [[("Fig. 4a - Algorithmic map: from the published objects to the findings "
                      "report", True, False)]], 24)
    set_text(s, 22, lead("Publishes in advance, versioned: ",
                         "the record protocol, the checks, the standard, the mapping from "
                         "findings to actions, the interest criteria."), 8)
    set_text(s, 24, lead("Review history: ",
                         "the journal witnesses the record from first submission through every "
                         "review round, so nothing can be backdated; it records what changed and "
                         "who changed it, author or AI."), 10)
    set_text(s, 63, ["review history, from first submission"], 8)
    shape_by_id(s, 63).width = Inches(2.2)
    set_text(s, 27, lead("The author works however they like: ",
                         "AI keeps the record to the protocol as the work proceeds. The working "
                         "history accrues from the start."), 8)
    set_text(s, 29, lead("The research record: ",
                         "the narrative and the AI appendix; materials; provenance linking each "
                         "result to what produced it; declarations (type, standard, AI use)."), 8)
    set_text(s, 31, lead("Submission: ",
                         "the record, with the author's own findings report. Working history at "
                         "the author's choice."), 8)
    set_text(s, 34, lead("Self-check: ",
                         "the author runs the same checks first and attaches the report; the "
                         "journal compares its run to it."), 7.5)
    set_text(s, 36, lead("Level 1 · Form: ",
                         "does the record conform to the published protocol for its declared type? "
                         "Findings; nothing is rejected on form. Internal."), 9)
    set_text(s, 38, lead("Level 2 · Correspondence: ",
                         "does the narrative match the AI appendix, and the appendix the record? "
                         "Internal: integrity."), 9)
    set_text(s, 42, [[("Level 3 · Warrant (SIEVE)", True, False)]], 9, align=PP_ALIGN.CENTER)
    set_text(s, 49, lead("Findings report: ",
                         "each finding with a pointer to the record and to the standard applied, "
                         "versioned with the model that produced it; plus the evidence map. Never "
                         "a verdict."), 9)
    set_text(s, 52, ["Checks differ by modality: formal built for Nov 1; empirical in Stage 2; "
                     "simulation, qualitative, verbal theory, and survey/review specified. All "
                     "three levels run on every submission, and the author runs them first."], 9,
             align=PP_ALIGN.CENTER)
    # Level 3 modules: contribution (the evidence map) is built; the rest specified
    for sid, label in ((43, "contribution"), (45, "internal"), (47, "soundness"),
                       (44, "statistical"), (46, "construct"), (48, "external")):
        set_text(s, sid, [[(label, False, False)]], 7, align=PP_ALIGN.CENTER)
        shape_by_id(s, sid).line.dash_style = MSO_LINE.SOLID
        set_fill(s, sid, "tint" if sid == 43 else "white", role="machine")
    set_fill(s, 24, "white")                # review history: specified
    set_fill(s, 34, "tint", role="author")  # self-check: the same checks, built
    set_text(s, 54, lead("Intake decision: ",
                         "a named editor applies the published mapping to the findings report "
                         "(Fig. 4b)."), 9, align=PP_ALIGN.CENTER)
    delete(s, 25, 32, 50, 84, 85, 86)
    for sid in (24, 31, 49):
        shape_by_id(s, sid).line.dash_style = MSO_LINE.SOLID


def fig4b_map(prs):
    """Dimitrija's Fig. 2b (slide 4) → Fig. 4b: findings report to certificate, per appendix
    C.3 / F / G. The evidence map moves left to feed the 'new' signature before review."""
    s = prs.slides[3]
    set_text(s, 2, [[("Fig. 4b - Algorithmic map: from the findings report to the certificate",
                      True, False)]], 24)
    heads = {15: "From Fig. 4a: the findings", 16: "Depart from the mapping?",
             17: "Standard action, signed", 18: "New, then review", 19: "Interest",
             20: "Certify", 21: "Outcomes"}
    for sid, h in heads.items():
        set_text(s, sid, [[(h, False, True)]], 9, align=PP_ALIGN.CENTER, color=MUTED)
    set_text(s, 22, lead("Review history: ",
                         "the journal witnesses the record from first submission through every "
                         "review round, so nothing can be backdated; it records what changed and "
                         "who changed it, author or AI. It stays with the journal."), 10)
    set_text(s, 25, lead("Certificate: ",
                         "standard, check, and model versions; data tier; new, in this kind, "
                         "with the map; of interest to this community; signed by three named "
                         "people."), 8)
    set_text(s, 27, lead("Back to the author: ",
                         "corrections (record fixed, no new work, checks re-run) or revise and "
                         "resubmit (new work; each round a new submission against the same review "
                         "history, Fig. 4a)."), 8)
    set_text(s, 29, lead("Findings report (Fig. 4a): ",
                         "findings with pointers, and the evidence map. No decision."), 8)
    set_text(s, 31, lead("Evidence map: ",
                         "the contribution claim against the literature the corpus covers: "
                         "nearest prior work, what the paper adds, of which kind."), 8)
    set_text(s, 33, lead("Editor reads the findings: ",
                         "a named editor applies the published mapping to the findings report."), 9)
    set_text(s, 35, ["Does the editor depart from the standard action?"], 8.5, align=PP_ALIGN.CENTER)
    set_text(s, 37, lead("Standard action, signed: ",
                         "proceed to review; return for corrections; return without review "
                         "(strengthen the evidence or weaken the claims); refer (a specialist "
                         "applies the mapping, then back to this step)."), 8)
    set_text(s, 39, lead("New, signed: ",
                         "one editor, on the evidence map, after the author answers it once; not "
                         "new → returned with the map. Then review rounds: reviewers receive the "
                         "narrative, findings, and map; they judge, and do not verify."), 8)
    set_text(s, 41, lead("Interest, signed: ",
                         "the editor and two named reviewers judge whether what is new is of "
                         "interest to this community, under the published criteria."), 8.5)
    set_text(s, 44, ["Certify: the editor issues the one certificate, or nothing."], 8,
             align=PP_ALIGN.CENTER)
    set_text(s, 46, lead("Returned: ",
                         "the record goes back unchanged. Findings, map, and review history stay "
                         "with the journal and travel only at the author's request."), 8.5)
    set_text(s, 49, ["no → standard action"], 8)
    set_text(s, 50, ["yes → the reason is recorded in the review history; the departure rate is "
                     "published"], 8)
    set_text(s, 56, ["return (corrections, or without review)"], 8)
    set_text(s, 59, ["the map: input, not judgment"], 8)
    set_text(s, 62, ["revise"], 8)
    set_text(s, 65, ["certify"], 8)
    set_text(s, 88, ["nothing"], 8)
    # the evidence map feeds the 'new' signature: move it over the New/review column
    dx = Inches(6.37) - shape_by_id(s, 31).left
    for sid in (31, 32, 59):
        sh = shape_by_id(s, sid); sh.left = sh.left + dx
    delete(s, 23, 42, 47, 58, 80, 81, 82)
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.12), Inches(4.75),
                                Inches(7.12), Inches(4.94))
    ln.line.color.rgb = GREY_LINE; ln.line.width = Pt(1); ln.line.dash_style = MSO_LINE.ROUND_DOT
    for sid in (22, 41, 46):
        shape_by_id(s, sid).line.dash_style = MSO_LINE.SOLID
    for sid in (22, 27, 41, 46):            # review history, revision, interest, returned: specified
        set_fill(s, sid, "white")
    set_text(s, 32, ["M6"], 7, align=PP_ALIGN.CENTER)   # was E1
    set_text(s, 43, ["H3"], 7, align=PP_ALIGN.CENTER)   # was W


def legend_slide(prs):
    """Slide 1: the legend, to the decided vocabulary. Dimitrija's note on greyscale stays."""
    s = prs.slides[0]
    set_text(s, 9, lead("Journal ", "— sets the criteria in public; certifies."), 10)
    set_text(s, 11, lead("Author + AI ", "— works however they like; AI keeps the record."), 10)
    set_text(s, 13, lead("Machine ", "— verifies; emits findings with pointers, never a decision."), 10)
    set_text(s, 15, lead("Named people ", "— editors and reviewers; decide, and sign."), 10)
    set_text(s, 17, lead("Limit ", "— a stated limit of the design, with its response."), 10)
    set_text(s, 19, lead("Neutral ", "— open, or out of scope."), 10)
    set_text(s, 30, ["Today's arrangement (Fig. 1) – grey"], 10)
    set_text(s, 39, ["Findings report; the certificate"], 10)
    set_text(s, 33, ["Artifact or published object"], 10)
    set_text(s, 31, [[("Role is also carried by lane position and by shape, so colour is never "
                       "the only channel; the states survive a greyscale print.", False, True)]],
             9, color=MUTED)
    delete(s, 26, 27, 28)


def reorder(prs, order):
    """Reorder slides by current index."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for i in order:
        lst.append(ids[i])


def main():
    import shutil
    shutil.copyfile(HERE / "system_diagrams v2claude.pptx", DECK)  # idempotent rebuild
    prs = Presentation(DECK)
    fig2_path(prs)      # slide 1 → Fig. 4
    fig3_levels(prs)    # slide 4 → Fig. 3
    fig4_record(prs)    # slide 5 → Fig. 2
    fig4a_map(prs)      # slide 2 → Fig. 4a
    fig4b_map(prs)      # slide 3 → Fig. 4b
    legend_slide(prs)   # slide 0
    fig1_before_after(prs)  # appended as slide 6 → Fig. 1
    reorder(prs, [0, 6, 5, 4, 1, 2, 3])  # legend, Fig. 1, 2, 3, 4, 4a, 4b
    prs.save(DECK)
    print("saved", DECK, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
