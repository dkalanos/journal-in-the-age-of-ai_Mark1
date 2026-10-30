"""Build system_diagrams v7dk.pptx from a fresh copy of v6dk.

Run:  python figures_v7.py
Then: powershell -NoProfile -File export_figures.ps1
      (PowerPoint; writes PNG and PDF per slide into this folder)

What changes from v6dk (Kalanoski, 2026-10-30), after a check of the legend against the figures:
- The "stretch goal" state (white fill, dashed border) is removed from the legend and from the key strips of
  Figs. 2, 4, 4a and 4b: no box in any figure is drawn that way.
- The legend's Limit swatch is drawn as the figures draw limits, white fill with an amber border, and a "limit"
  entry is added to the key strips of Figs. 2, 3 and 4, the figures that use amber-bordered boxes.
- Fig. 1's "Today" column, its arrows included, uses the legend's Neutral grey (theme accent 6) in place of two
  greys set by hand, and the legend's Neutral swatch takes the same light tint as the other role swatches.
- Fig. 4b's "New, signed" (H5) is tinted like the other signed steps: signing "new" is built in Mark 1; the
  demonstration stops before it, which is not a state the legend has.
- Fig. 2's "What is new, taken together" box takes the Neutral border in place of the Journal one: it is a note, and
  white fill with a role border is the legend's "specified, not built", which its five items are not.
- Fig. 4a's M2 "Checks differ by modality" is tinted: the router that reads the declared kind of work and sends the
  formal kind to its checks runs in Mark 1 (machine/modality.py, "built; formal only" in its README); the box's text
  says which kinds are built. M4 is drawn the same way, tinted with its unbuilt modules white inside it.
Everything else is as in v6dk. Shape ids below are v6dk's.
"""
import copy
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches

HERE = Path(__file__).parent
SRC = HERE / "system_diagrams v6dk.pptx"
DECK = HERE / "system_diagrams v7dk.pptx"
TINT = 0.78  # the light fill of every role box: lumMod 22000 / lumOff 78000


def by_id(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(f"slide has no shape {sid}")


def tint(shape, theme, brightness=TINT):
    shape.fill.solid()
    shape.fill.fore_color.theme_color = theme
    if brightness:
        shape.fill.fore_color.brightness = brightness


def white(shape):
    shape.fill.solid()
    shape.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
    shape.fill.fore_color.brightness = 0  # drop the old tint, so the swatch is encoded as the figures' limit boxes are


def border(shape, theme):
    shape.line.color.theme_color = theme


def delete(slide, *sids):
    for sid in sids:
        el = by_id(slide, sid)._element
        el.getparent().remove(el)


def clone(slide, sid, left, text=None):
    """Copy shape `sid` to a new x position (same y), with a fresh id and, for a text box, new text."""
    src = by_id(slide, sid)
    new = copy.deepcopy(src._element)
    new_id = max(s.shape_id for s in slide.shapes) + 1
    new.xpath("./p:nvSpPr/p:cNvPr")[0].set("id", str(new_id))
    slide.shapes._spTree.append(new)
    shape = by_id(slide, new_id)
    shape.left = left
    if text is not None:
        runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    return shape


def legend(prs):
    s = prs.slides[0]
    white(by_id(s, 16))                                   # Limit: as the figures draw it
    tint(by_id(s, 18), MSO_THEME_COLOR.ACCENT_6)          # Neutral: the light tint the other roles use
    delete(s, 24, 25)                                     # stretch goal, swatch and text


def key_strip(slide, swatch_id, text_id, specified_swatch_id, specified_text_id, left):
    """Replace a 'stretch' entry (or add one, when ids are None) with a 'limit' entry at `left`."""
    if swatch_id is not None:
        delete(slide, swatch_id, text_id)
    sw = clone(slide, specified_swatch_id, left)
    border(sw, MSO_THEME_COLOR.ACCENT_5)
    clone(slide, specified_text_id, left + Inches(0.40), "limit")


def fig1(prs):
    s = prs.slides[1]
    for sid in (8, 12, 17, 22, 27, 32, 36, 40):           # the "Today" column
        tint(by_id(s, sid), MSO_THEME_COLOR.ACCENT_6)
        border(by_id(s, sid), MSO_THEME_COLOR.ACCENT_6)
    for sid in (9, 14, 19, 24, 29, 33, 37, 41):           # its arrows
        tint(by_id(s, sid), MSO_THEME_COLOR.ACCENT_6, brightness=0)


def fig2(prs):
    s = prs.slides[2]
    border(by_id(s, 28), MSO_THEME_COLOR.ACCENT_6)        # "What is new, taken together": a note, not a specified object


def fig4a(prs):
    s = prs.slides[5]
    tint(by_id(s, 52), MSO_THEME_COLOR.ACCENT_3)          # M2, the modality router: built, formal only


def fig4b(prs):
    s = prs.slides[6]
    tint(by_id(s, 39), MSO_THEME_COLOR.ACCENT_4)          # H5 "New, signed": built


def main():
    shutil.copyfile(SRC, DECK)  # idempotent rebuild from v6dk
    prs = Presentation(DECK)
    legend(prs)
    fig1(prs)
    fig2(prs)
    key_strip(prs.slides[2], 41, 42, 39, 40, Inches(7.24))      # Fig. 2
    key_strip(prs.slides[3], None, None, 58, 59, Inches(7.25))  # Fig. 3 (had no stretch entry)
    key_strip(prs.slides[4], 53, 54, 51, 52, Inches(7.27))      # Fig. 4
    delete(prs.slides[5], 82, 83)                               # Fig. 4a: stretch only
    fig4a(prs)
    delete(prs.slides[6], 78, 79)                               # Fig. 4b: stretch only
    fig4b(prs)
    prs.save(DECK)
    print("saved", DECK, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
